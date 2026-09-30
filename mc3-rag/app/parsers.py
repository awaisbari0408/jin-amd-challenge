"""Turn every readable file in the corpus into text records.

Rules the graded corpus enforces:
- One bad file must never stop the walk: every file is parsed inside its own
  try/except, and failures are recorded instead of raised.
- Unreadable (chmod 000), encrypted and unknown-type files are skipped; nothing
  inside them may ever be an answer.
- File type is decided by EXTENSION only. Some junk files look like images to
  content sniffers.
"""
import csv
import os
import re

TEXT_EXT = {".txt", ".log", ".py", ".md", ".json", ".yaml", ".yml", ".ini", ".cfg", ".toml"}
IMAGE_EXT = {".png", ".jpg", ".jpeg"}
SUPPORTED = TEXT_EXT | IMAGE_EXT | {".pdf", ".docx", ".xlsx", ".csv"}

_SUPERSEDED = re.compile(r"\b(withdrawn|superseded by|obsolete|deprecated|do not use)\b", re.I)


def walk(corpus):
    """Yield corpus-relative paths (forward slashes), sorted, never raising."""
    for root, dirs, files in os.walk(corpus, onerror=lambda e: None):
        dirs.sort()
        for name in sorted(files):
            full = os.path.join(root, name)
            yield os.path.relpath(full, corpus).replace(os.sep, "/")


def _pdf(path):
    import pypdf
    reader = pypdf.PdfReader(path)
    if reader.is_encrypted:  # must be checked before touching pages
        raise PermissionError("encrypted PDF")
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def _docx(path):
    import docx
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if not rows:
            continue
        head = rows[0]
        for r in rows[1:]:
            out.append(" | ".join(f"{h}: {v}" if h else v for h, v in zip(head, r)))
    return "\n".join(out)


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [tuple("" if v is None else str(v) for v in r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(x.strip() for x in r)]
        if not rows:
            continue
        head = rows[0]
        if len(rows) == 1 or len([h for h in head if h]) < 2:
            out.extend(f"[sheet {ws.title}] " + " ".join(r) for r in rows)
            continue
        for r in rows[1:]:
            cells = " | ".join(f"{h}: {v}" for h, v in zip(head, r) if v)
            out.append(f"[sheet {ws.title}] {cells}")
    return "\n".join(out)


def _csv(path):
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        rows = list(csv.reader(f))
    rows = [r for r in rows if any(x.strip() for x in r)]
    if len(rows) < 2:
        return "\n".join(",".join(r) for r in rows)
    head = rows[0]
    return "\n".join(" | ".join(f"{h}: {v}" for h, v in zip(head, r) if v) for r in rows[1:])


def _text(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def parse_file(corpus, rel, read_image=None):
    """Return (text, note). text is None when the file is skipped."""
    path = os.path.join(corpus, rel)
    ext = os.path.splitext(rel)[1].lower()
    if ext not in SUPPORTED:
        return None, "unknown type"
    try:
        with open(path, "rb"):
            pass  # raises PermissionError for chmod 000 / dropped DAC_OVERRIDE
        if ext == ".pdf":
            return _pdf(path), "pdf"
        if ext == ".docx":
            return _docx(path), "docx"
        if ext == ".xlsx":
            return _xlsx(path), "xlsx"
        if ext == ".csv":
            return _csv(path), "csv"
        if ext in IMAGE_EXT:
            if read_image is None:
                return None, "image skipped (no vision model)"
            return read_image(path), "image transcription"
        return _text(path), "text"
    except Exception as e:  # noqa: BLE001 - one bad file must never stop the walk
        return None, f"skipped: {e.__class__.__name__}: {e}"[:200]


def is_superseded(rel, text):
    """A withdrawn/superseded document is never a valid source."""
    # Only the filename and the title lines: a VALID document may mention in its
    # revision history that an older revision "is withdrawn".
    head = "\n".join([l for l in (text or "").splitlines() if l.strip()][:3])
    return bool(_SUPERSEDED.search(rel.replace("_", " ")) or _SUPERSEDED.search(head))


def chunk(text, size=700, overlap=120):
    """Split on lines into ~size-char chunks; one table row / log line never splits."""
    lines = [l.rstrip() for l in (text or "").splitlines() if l.strip()]
    chunks, cur = [], ""
    for line in lines:
        if cur and len(cur) + len(line) + 1 > size:
            chunks.append(cur)
            cur = cur[-overlap:].split("\n", 1)[-1] if overlap else ""
        cur = f"{cur}\n{line}" if cur else line
    if cur:
        chunks.append(cur)
    return chunks
